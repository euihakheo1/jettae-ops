"""Shared helpers for the ``test_api_*`` suites (no tests here).

* :func:`settings` / :func:`new_service` / :func:`new_app` — an API on an in-memory store with a
  fixed clock and the LLM pinned OFF, independent of the developer's ``JETTAE_*`` environment.
* :func:`open_client` — a ``TestClient`` whose lifespan is entered once per block. Sandboxed
  Windows runners intermittently fail to create an event loop's self-pipe (WinError 10014), which
  has nothing to do with the code under test, so entering is retried.
* Role headers, the demo bundle files, and in-test builders for PDF/DOCX/HTML runbooks.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import docx
from fastapi import FastAPI
from fastapi.testclient import TestClient

from jettae.api import create_app
from jettae.config import Settings
from jettae.domain import LlmMode
from jettae.orchestrator import JettaeService
from jettae.store import MemoryStore, Store

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "examples" / "dmc-shop"
API = "/api/v1"
CLOCK = datetime(2026, 10, 28, 14, 0, tzinfo=UTC)

VIEWER = {"X-Jettae-User": "demo.viewer", "X-Jettae-Role": "viewer"}
EDITOR = {"X-Jettae-User": "demo.editor", "X-Jettae-Role": "editor"}
LEAD = {"X-Jettae-User": "park.jaehyun", "X-Jettae-Role": "lead"}
ADMIN = {"X-Jettae-User": "demo.admin", "X-Jettae-Role": "admin"}

#: A tiny runbook the rule parser understands, for tests that need a quick real diagnosis.
SMALL_RUNBOOK = (
    "# 결제 서비스 런북\n\n## 결제 서비스\n\n"
    "- 감지: 결제 실패율 알람 → 김서연 사원 호출\n"
    "- 담당: 김서연 사원 (슬랙, 전화)\n"
)
SMALL_ROSTER = "이름,상태,직급,팀,역할,별칭,연락수단,퇴사일\n김서연,재직,사원,결제팀,1차 당직,,슬랙;전화,\n"

_PINNED: dict[str, Any] = {
    "llm_mode": LlmMode.OFF,
    "store_raw_documents": False,
    "database_url": "memory://",
    "web_dist": None,
    "auth_mode": "dev",
    "api_tokens": {},
}
_ENTER_ATTEMPTS = 8


def fixed_clock() -> datetime:
    return CLOCK


def settings(**overrides: Any) -> Settings:
    """Settings independent of the developer's environment and ``.env`` file."""
    return Settings(_env_file=None, **{**_PINNED, **overrides})


def new_service(
    *, store: Store | None = None, clock: Callable[[], datetime] = fixed_clock, **overrides: Any
) -> JettaeService:
    target = store if store is not None else MemoryStore(clock=clock)
    return JettaeService(target, settings(**overrides), clock=clock)


def new_app(
    *, service: JettaeService | None = None, clock: Callable[[], datetime] = fixed_clock, **overrides: Any
) -> FastAPI:
    """An app on its own in-memory service (``overrides`` go to the API's settings)."""
    api_settings = settings(**overrides)
    return create_app(api_settings, service=service if service is not None else new_service(clock=clock))


@contextmanager
def open_client(app: FastAPI, **kwargs: Any) -> Iterator[TestClient]:
    """``with TestClient(app)`` (lifespan + one event loop for the block), retried on loop creation."""
    for attempt in range(_ENTER_ATTEMPTS):
        client = TestClient(app, **kwargs)
        try:
            client.__enter__()
        except OSError:
            if attempt == _ENTER_ATTEMPTS - 1:
                raise
            continue
        break
    try:
        yield client
    finally:
        client.__exit__(None, None, None)


def demo_files() -> list[tuple[str, tuple[str, bytes, str]]]:
    """The demo bundle as multipart parts (runbooks + structured sources + drills)."""
    parts: list[tuple[str, tuple[str, bytes, str]]] = [
        ("runbooks", (path.name, path.read_bytes(), "text/markdown"))
        for path in sorted((DEMO / "runbooks").iterdir())
    ]
    for field, name in (
        ("roster", "roster.csv"),
        ("oncall", "oncall.csv"),
        ("catalog", "catalog.yaml"),
        ("policy", "policy.yaml"),
        ("scenarios", "scenarios.yaml"),
    ):
        parts.append((field, (name, (DEMO / name).read_bytes(), "application/octet-stream")))
    return parts


def create_workspace(client: TestClient, name: str = "결제 플랫폼팀") -> str:
    response = client.post(f"{API}/workspaces", json={"name": name}, headers=EDITOR)
    assert response.status_code == 201, response.text
    workspace_id: str = response.json()["id"]
    return workspace_id


# --------------------------------------------------------------------------- document builders


def pdf_bytes(lines: list[str]) -> bytes:
    """A minimal valid one-page PDF (Helvetica, ASCII text), one text line per entry."""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [4 0 R] /Count 1 >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
        b"/Contents 5 0 R >>",
    ]
    ops = ["BT", "/F1 12 Tf", "72 720 Td"]
    for index, line in enumerate(lines):
        if index:
            ops.append("0 -16 Td")
        ops.append(f"({line}) Tj")
    ops.append("ET")
    stream = "\n".join(ops).encode()
    objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def docx_bytes(heading: str, paragraphs: list[str]) -> bytes:
    document = docx.Document()
    document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def html_bytes(heading: str, items: list[str]) -> bytes:
    body = "".join(f"<li>{item}</li>" for item in items)
    return f"<!doctype html><html><body><h1>{heading}</h1><ul>{body}</ul></body></html>".encode()
